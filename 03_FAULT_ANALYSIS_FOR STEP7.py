#!/usr/bin/env python3
"""Export ECFR fault-event Parquet data to a metrics-focused Excel workbook.

The workbook emphasizes fault codes and event counts while retaining optional
row-level data sheets. It streams the Parquet file and supports Excel's row cap.
"""
from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DEFAULT_ROOT = Path('/raid3/e296408/All_ECFRs_working/ECFR_11001_ALL_FAULTS')
DEFAULT_INPUT = DEFAULT_ROOT / 'all_ecfr_fault_events.parquet'
DEFAULT_OUTPUT = DEFAULT_ROOT / 'all_ecfr_fault_metrics.xlsx'
ROWS_PER_DATA_SHEET = 900_000

NAVY = '07111F'
PANEL = '0D2033'
WHITE = 'F3F7FA'
GREEN = '008000'
TEAL = '0D9488'
AMBER = 'F59E0B'
LIGHT_TEAL = 'DDF6F1'
LIGHT_AMBER = 'FFF0D5'
GRAY = '666666'


def excel_value(value):
    if value is None:
        return None
    if isinstance(value, (list, tuple, set, np.ndarray)):
        return json.dumps(list(value), default=str, ensure_ascii=False)
    if isinstance(value, dict):
        return json.dumps(value, default=str, ensure_ascii=False)
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        if value.tzinfo is not None:
            value = value.tz_convert('UTC').tz_localize(None)
        return value.to_pydatetime()
    return value


def safe_table_name(workbook, suffix):
    safe = re.sub(r'[^A-Za-z0-9_]', '', str(suffix)) or 'Table'
    base = f'Fault_{safe}'
    existing = {
        table.displayName
        for sheet in workbook.worksheets
        for table in sheet.tables.values()
    }
    name = base[:240]
    counter = 2
    while name in existing:
        tail = f'_{counter}'
        name = base[:240-len(tail)] + tail
        counter += 1
    return name


def add_table(ws, first_row, last_row, last_col, suffix):
    if last_row <= first_row or last_col <= 0:
        return
    ref = f'A{first_row}:{get_column_letter(last_col)}{last_row}'
    table = Table(displayName=safe_table_name(ws.parent, suffix), ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium2',
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def title(ws, text, subtitle=None, max_col=10):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.cell(1, 1, text)
    ws.cell(1, 1).fill = PatternFill('solid', fgColor=NAVY)
    ws.cell(1, 1).font = Font(color=WHITE, bold=True, size=18)
    ws.cell(1, 1).alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        ws.cell(2, 1, subtitle)
        ws.cell(2, 1).font = Font(color=GRAY, italic=True, size=10)
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[2].height = 27


def header(ws, row, labels):
    for index, label in enumerate(labels, 1):
        cell = ws.cell(row, index, label)
        cell.fill = PatternFill('solid', fgColor=PANEL)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 30


def write_dataframe(ws, frame, start_row=4, table_suffix=None):
    labels = list(frame.columns)
    header(ws, start_row, labels)
    for row_number, values in enumerate(frame.itertuples(index=False, name=None), start_row + 1):
        for column_number, value in enumerate(values, 1):
            ws.cell(row_number, column_number, excel_value(value))
    if table_suffix and not frame.empty:
        add_table(ws, start_row, start_row + len(frame), len(labels), table_suffix)
    return start_row + len(frame)


def choose_column(columns, candidates):
    lookup = {column.lower(): column for column in columns}
    for candidate in candidates:
        if candidate.lower() in lookup:
            return lookup[candidate.lower()]
    return None


def build_metrics(input_path):
    parquet = pq.ParquetFile(input_path)
    columns = parquet.schema_arrow.names
    code_col = choose_column(columns, ['fault_code_normalized', 'fault_code', 'fault_codes'])
    if not code_col:
        raise ValueError('No fault-code column found in the Parquet data')

    engine_col = choose_column(columns, ['engine_serial', 'engine_key'])
    hash_col = choose_column(columns, ['txt_sha256', 'file_hash_partition'])
    file_col = choose_column(columns, ['source_txt_file'])
    time_col = choose_column(columns, ['record_datetime_utc', 'first_fault_time'])
    leg_col = choose_column(columns, ['leg_number', 'resolved_leg'])

    selected = [code_col]
    for column in [engine_col, hash_col, file_col, time_col, leg_col]:
        if column and column not in selected:
            selected.append(column)

    code_counter = {}
    code_engines = {}
    code_hashes = {}
    code_files = {}
    code_legs = {}
    code_first = {}
    code_last = {}
    monthly = {}
    total_rows = 0

    for batch in parquet.iter_batches(batch_size=100_000, columns=selected):
        frame = batch.to_pandas()
        frame[code_col] = frame[code_col].fillna('').astype(str).str.strip()
        frame = frame.loc[frame[code_col].ne('')].copy()
        total_rows += len(frame)

        if time_col:
            frame[time_col] = pd.to_datetime(frame[time_col], errors='coerce', utc=True)

        for code, group in frame.groupby(code_col, dropna=False):
            code = str(code)
            code_counter[code] = code_counter.get(code, 0) + len(group)
            if engine_col:
                code_engines.setdefault(code, set()).update(
                    group[engine_col].dropna().astype(str).str.strip().loc[lambda x: x.ne('')]
                )
            if hash_col:
                code_hashes.setdefault(code, set()).update(
                    group[hash_col].dropna().astype(str).str.strip().loc[lambda x: x.ne('')]
                )
            if file_col:
                code_files.setdefault(code, set()).update(
                    group[file_col].dropna().astype(str).str.strip().loc[lambda x: x.ne('')]
                )
            if leg_col:
                code_legs.setdefault(code, set()).update(group[leg_col].dropna().astype(str))
            if time_col:
                valid_times = group[time_col].dropna()
                if not valid_times.empty:
                    candidate_first = valid_times.min()
                    candidate_last = valid_times.max()
                    code_first[code] = min(code_first.get(code, candidate_first), candidate_first)
                    code_last[code] = max(code_last.get(code, candidate_last), candidate_last)

        if time_col:
            dated = frame.loc[frame[time_col].notna()].copy()
            if not dated.empty:
                dated['month'] = dated[time_col].dt.tz_convert('UTC').dt.tz_localize(None).dt.to_period('M').astype(str)
                counts = dated.groupby('month').size()
                for month, count in counts.items():
                    monthly[month] = monthly.get(month, 0) + int(count)

    rows = []
    for code, count in code_counter.items():
        rows.append({
            'Fault Code': code,
            'Event Count': int(count),
            'Affected Engines': len(code_engines.get(code, set())),
            'Source Hashes': len(code_hashes.get(code, set())),
            'Source Files': len(code_files.get(code, set())),
            'Distinct Legs': len(code_legs.get(code, set())),
            'First Observed UTC': code_first.get(code),
            'Last Observed UTC': code_last.get(code),
        })

    summary = pd.DataFrame(rows).sort_values(['Event Count', 'Fault Code'], ascending=[False, True]).reset_index(drop=True)
    summary.insert(0, 'Rank', np.arange(1, len(summary) + 1))
    summary['Percent of Events'] = summary['Event Count'] / max(1, summary['Event Count'].sum())
    summary['Cumulative Percent'] = summary['Percent of Events'].cumsum()
    summary['Events per Engine'] = summary['Event Count'] / summary['Affected Engines'].replace(0, np.nan)
    summary['Events per Source Hash'] = summary['Event Count'] / summary['Source Hashes'].replace(0, np.nan)

    monthly_frame = pd.DataFrame(
        [{'Month': month, 'Event Count': count} for month, count in sorted(monthly.items())]
    )
    return parquet, summary, monthly_frame, {
        'fault_code_column': code_col,
        'engine_column': engine_col,
        'hash_column': hash_col,
        'file_column': file_col,
        'time_column': time_col,
        'leg_column': leg_col,
        'fault_rows': int(total_rows),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', type=Path, default=DEFAULT_INPUT)
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument('--include-detail', action='store_true')
    parser.add_argument('--rows-per-sheet', type=int, default=ROWS_PER_DATA_SHEET)
    args = parser.parse_args()

    if not args.input.is_file():
        raise FileNotFoundError(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)

    parquet, metrics, monthly, detected = build_metrics(args.input)
    total_events = int(metrics['Event Count'].sum())
    total_codes = int(len(metrics))
    top_10_share = float(metrics.head(10)['Percent of Events'].sum()) if total_codes else 0
    codes_80 = int((metrics['Cumulative Percent'].lt(0.8)).sum() + 1) if total_codes else 0

    wb = Workbook()
    wb.remove(wb.active)

    # Dashboard
    ws = wb.create_sheet('Dashboard')
    title(ws, 'All ECFR Fault Metrics Dashboard', f'Source: {args.input}', 12)
    kpis = [
        ('Total fault events', total_events),
        ('Distinct fault codes', total_codes),
        ('Top 10 event share', top_10_share),
        ('Codes covering 80%', codes_80),
        ('Affected engines', int(metrics['Affected Engines'].max()) if total_codes else 0),
        ('Source hashes', int(metrics['Source Hashes'].max()) if total_codes else 0),
    ]
    for i, (label, value) in enumerate(kpis):
        column = 1 + (i % 3) * 4
        row = 4 + (i // 3) * 3
        ws.merge_cells(start_row=row, start_column=column, end_row=row, end_column=column + 2)
        ws.cell(row, column, label)
        ws.cell(row, column).fill = PatternFill('solid', fgColor=PANEL)
        ws.cell(row, column).font = Font(color=WHITE, bold=True)
        ws.merge_cells(start_row=row + 1, start_column=column, end_row=row + 1, end_column=column + 2)
        ws.cell(row + 1, column, value)
        ws.cell(row + 1, column).font = Font(color=TEAL, bold=True, size=22)
        ws.cell(row + 1, column).fill = PatternFill('solid', fgColor=LIGHT_TEAL)
        ws.cell(row + 1, column).alignment = Alignment(horizontal='center')
        if 'share' in label.lower():
            ws.cell(row + 1, column).number_format = '0.0%'

    top = metrics.head(20)[['Fault Code', 'Event Count', 'Percent of Events', 'Affected Engines']]
    write_dataframe(ws, top, 11, 'DashboardTop20')
    for row in range(12, 12 + len(top)):
        ws.cell(row, 3).number_format = '0.0%'

    chart = BarChart()
    chart.type = 'bar'
    chart.style = 10
    chart.title = 'Top 15 Fault Codes by Event Count'
    chart.y_axis.title = 'Fault Code'
    chart.x_axis.title = 'Events'
    chart.add_data(Reference(ws, min_col=2, min_row=11, max_row=26), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=12, max_row=26))
    chart.height = 8
    chart.width = 14
    chart.legend = None
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    ws.add_chart(chart, 'F11')
    for col, width in {'A':18, 'B':16, 'C':18, 'D':18, 'F':2}.items():
        ws.column_dimensions[col].width = width

    # Full fault-code metrics
    ws = wb.create_sheet('Fault Code Metrics')
    title(ws, 'Fault Code Metrics', 'Sorted by event count. Percentages are based on total extracted fault events.', 12)
    write_dataframe(ws, metrics, 4, 'FaultCodeMetrics')
    ws.freeze_panes = 'A5'
    for row in range(5, 5 + len(metrics)):
        ws.cell(row, 10).number_format = '0.0%'
        ws.cell(row, 11).number_format = '0.0%'
        ws.cell(row, 12).number_format = '0.00'
        ws.cell(row, 13).number_format = '0.00'
    widths = [8, 20, 16, 18, 18, 18, 16, 23, 23, 18, 19, 19, 22]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Monthly trend
    ws = wb.create_sheet('Monthly Trend')
    title(ws, 'Fault Events by Month', 'Monthly trend is included when a usable event timestamp exists.', 8)
    if not monthly.empty:
        write_dataframe(ws, monthly, 4, 'MonthlyTrend')
        chart = LineChart()
        chart.title = 'Monthly Fault Event Trend'
        chart.y_axis.title = 'Events'
        chart.x_axis.title = 'Month'
        chart.add_data(Reference(ws, min_col=2, min_row=4, max_row=4 + len(monthly)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(monthly)))
        chart.height = 10
        chart.width = 19
        chart.legend = None
        ws.add_chart(chart, 'D4')
    else:
        ws['A4'] = 'No usable timestamp column was found.'
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 18

    # Metadata
    ws = wb.create_sheet('Export Metadata')
    title(ws, 'Export Metadata', 'Detected source columns and workbook generation settings.', 6)
    metadata_rows = [
        ('Input Parquet', str(args.input)),
        ('Output workbook', str(args.output)),
        ('Fault event rows', total_events),
        ('Distinct fault codes', total_codes),
        ('Detail sheets included', args.include_detail),
        ('Detail rows per sheet', args.rows_per_sheet),
    ] + [(key, value) for key, value in detected.items()]
    header(ws, 4, ['Metric', 'Value'])
    for row, item in enumerate(metadata_rows, 5):
        ws.cell(row, 1, item[0])
        ws.cell(row, 2, excel_value(item[1]))
    add_table(ws, 4, 4 + len(metadata_rows), 2, 'ExportMetadata')
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 100

    # Optional detail data, split safely below Excel's row limit.
    if args.include_detail:
        columns = parquet.schema_arrow.names
        sheet_number = 0
        detail_ws = None
        rows_in_sheet = 0
        for batch in parquet.iter_batches(batch_size=50_000):
            frame = batch.to_pandas()
            for values in frame.itertuples(index=False, name=None):
                if detail_ws is None or rows_in_sheet >= args.rows_per_sheet:
                    if detail_ws is not None:
                        add_table(detail_ws, 3, 3 + rows_in_sheet, len(columns), f'Detail{sheet_number:03d}')
                    sheet_number += 1
                    rows_in_sheet = 0
                    detail_ws = wb.create_sheet(f'Fault Data {sheet_number:03d}')
                    title(detail_ws, f'All ECFR Fault Events | Part {sheet_number:03d}', None, len(columns))
                    header(detail_ws, 3, columns)
                    detail_ws.freeze_panes = 'A4'
                    for index, column in enumerate(columns, 1):
                        name = column.lower()
                        width = 55 if ('source' in name or 'payload' in name or 'value_raw' in name) else 22
                        detail_ws.column_dimensions[get_column_letter(index)].width = width
                detail_ws.append([excel_value(value) for value in values])
                rows_in_sheet += 1
        if detail_ws is not None:
            add_table(detail_ws, 3, 3 + rows_in_sheet, len(columns), f'Detail{sheet_number:03d}')

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.oddFooter.center.text = 'All ECFR Fault Metrics'
        sheet.oddFooter.right.text = 'Page &P of &N'

    wb.save(args.output)
    print(json.dumps({
        'input': str(args.input),
        'output': str(args.output),
        'total_fault_events': total_events,
        'distinct_fault_codes': total_codes,
        'top_10_share': top_10_share,
        'codes_covering_80_percent': codes_80,
        'detail_included': args.include_detail,
    }, indent=2))


if __name__ == '__main__':
    main()
